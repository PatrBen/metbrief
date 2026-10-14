"""Build a plain PPTX (one slide per Skysight image, 16:9) from the images skysight_shots.py produced."""
import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt

MAX_EDGE = 1800   # px; keeps the PPTX small (full-size screenshots are ~5 MB each)


def build_pptx(targets, entries, base_dir, spot_name, day, out_file):
    """targets: from skysight_shots.targets(); entries: result list (only ok ones are placed)."""
    ok_files = {e['file'] for e in entries if e['ok']}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    props = prs.core_properties            # python-pptx's template carries a foreign author: neutralise
    props.author = props.last_modified_by = props.comments = props.keywords = ''
    props.title = f'Skysight {spot_name} {day.isoformat()}'
    blank = prs.slide_layouts[6]
    margin, top = Emu(150000), Emu(620000)
    area_w, area_h = prs.slide_width - 2 * margin, prs.slide_height - top - Emu(100000)
    for target in targets:
        if target['file'] not in ok_files:
            continue
        slide = prs.slides.add_slide(blank)
        title = slide.shapes.add_textbox(margin, Emu(80000), area_w, Emu(480000)).text_frame
        title.text = f"{target['title']} - {spot_name}, {day.isoformat()}"
        title.paragraphs[0].runs[0].font.size = Pt(24)
        if target.get('note'):
            slide.notes_slide.notes_text_frame.text = target['note']
        with Image.open(Path(base_dir) / target['file']) as img:
            img = img.convert('RGB')
            img.thumbnail((MAX_EDGE, MAX_EDGE))
            buf = io.BytesIO()
            img.save(buf, 'JPEG', quality=85)
            scale = min(area_w / img.width, area_h / img.height)
            width, height = int(img.width * scale), int(img.height * scale)
        slide.shapes.add_picture(buf, margin + (area_w - width) // 2, top + (area_h - height) // 2, width, height)
    prs.save(str(out_file))
